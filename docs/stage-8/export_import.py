# Шаг 8.4: выгрузка «Финансы 25.xlsm» и «Имущество КФХ.xlsx» в JSON для
# импортёра (packages/sync/tool/import_finance.dart). Только чтение Excel.
#
#   PYTHONUTF8=1 python docs/stage-8/export_import.py
#
# Результат — private/finance25_import.json (вне git: в нём данные владельца).
# uuid каждой записи вычисляется из её источника (uuid5): повторная выгрузка
# даёт те же uuid, повторный импорт не плодит дублей. Решения владельца — в
# DEVELOPMENT_PLAN.md (шаги 8.0, 8.4).

import collections
import datetime
import decimal
import json
import os
import sys
import uuid
import warnings

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_mapping as bm  # noqa: E402 — разбор «Финансов 25» и сопоставление

warnings.filterwarnings("ignore")
ASSETS_SRC = r"C:/Users/Иван/Yandex.Disk/КФХ/Имущество КФХ.xlsx"
TABEL_SRC = r"C:/Users/Иван/OneDrive/ЛИКА/Табель 25.xlsx"
NS = uuid.uuid5(uuid.NAMESPACE_URL, "https://tab.korovatech.ru/finance25")


def uid(*parts):
    return str(uuid.uuid5(NS, "/".join(str(p) for p in parts)))


def kop(value):
    """Рубли → копейки, половина — от нуля (как toKopecks в kfh_domain)."""
    d = decimal.Decimal(str(value)).quantize(decimal.Decimal("0.01"), rounding=decimal.ROUND_HALF_UP)
    return int(d * 100)


def day(value):
    return value.strftime("%Y-%m-%d")


def text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    t = str(value).replace("_x000D_", "").strip()
    return t or None


# --- Направления и ключ ----------------------------------------------------
DIR_NAMES = {"С": "Животноводство", "К": "Растениеводство (корма)", "Р": "Реализация", "О": "Общехозяйственные"}
directions = [
    {"id": uid("direction", code), "name": name, "overhead": code == "О", "sortOrder": i}
    for i, (code, name) in enumerate(DIR_NAMES.items())
]
dir_id = {code: uid("direction", code) for code in DIR_NAMES}
KEY_ID = uid("key", "Трактора")
keys = [{
    "id": KEY_ID,
    "name": "Трактора",
    "periods": [{"validFrom": "2025-01-01", "shares": {dir_id["С"]: 60, dir_id["К"]: 40}}],
}]


def target(code):
    """Код направления из сопоставления → (direction, key)."""
    return (None, KEY_ID) if code == "Т" else (dir_id[code], None)


# --- Статьи ----------------------------------------------------------------
KIND = {"тек": "current", "кап": "capex", "возврат": "refund"}
TYPE = {"расход": "expense", "доход": "income"}
categories = []
cat_id = {}
EXTRA = [
    ("доход", "3. Прочие поступления", "Корректировка остатка (излишек)", "О", "тек"),
    ("расход", "08. Управление", "Корректировка остатка (недостача)", "О", "тек"),
]
for i, (kind, group, name, code, vid) in enumerate(list(bm.cats) + EXTRA):
    direction, key = target(code)
    cid = uid("category", kind, name)
    cat_id[(kind, name)] = cid
    categories.append({
        "id": cid, "type": TYPE[kind], "group": group, "name": name,
        "directionId": direction, "allocationKeyId": key, "kind": KIND[vid],
        "inTaxBook": True, "hidden": False, "sortOrder": i,
    })
cat_code = {(c[0], c[2]): c[3] for c in bm.cats}

# --- Кошельки (решения владельца 2026-10-08) -------------------------------
WALLETS = [
    # имя в Excel, имя в программе, вид, сотрудник (подотчёт)
    ("Касса нал", "Касса нал", "cash", None),
    ("СБЕР", "СБЕР", "bank", None),
    ("РСХБ", "РСХБ", "bank", None),
    ("Карта ГКФХ", "Карта ГКФХ", "card", None),
    ("СБЕР Кредитка", "СБЕР Кредитка", "credit_card", None),
]
# Подотчёт людей: «имя в Excel;имя в программе;сотрудник табеля» —
# private/accountable.csv (вне git: имена людей).
people_file = os.path.join(bm.PRIVATE, "accountable.csv")
for line in open(people_file, encoding="utf-8").read().splitlines()[1:]:
    if line.count(";") == 2:
        excel, name, employee = (p.strip() for p in line.split(";"))
        WALLETS.insert(len(WALLETS) - 1, (excel, name, "accountable", employee))
opening = {bm.s(w["Счета"]): w["Начальный остаток"] for w in bm.wl}
accounts = []
acc_id = {}
for i, (excel, name, kind, employee) in enumerate(WALLETS):
    aid = uid("account", excel)
    acc_id[excel] = aid
    start = opening.get(excel)
    accounts.append({
        "id": aid, "name": name, "kind": kind,
        "openingBalance": kop(start) if start else 0,
        "openingDate": "2025-01-01" if start else None,
        "employeeName": employee, "closed": False, "sortOrder": i,
    })

# --- Контрагенты -----------------------------------------------------------
merge = {}
merge_file = os.path.join(bm.PRIVATE, "counterparty_merge.csv")
if os.path.exists(merge_file):
    for line in open(merge_file, encoding="utf-8").read().splitlines()[1:]:
        if ";" in line:
            a, b = line.split(";", 1)
            merge[a.strip()] = b.strip()


def cp_name(name):
    return merge.get(name, name)


seen = collections.defaultdict(set)
staff = set()
for r in bm.rows:
    if r["action"] != "статья" or not r["who"]:
        continue
    seen[cp_name(r["who"])].add(r["kind"])
    if r["old"].startswith("03."):
        staff.add(cp_name(r["who"]))
debt_names = [bm.s(d["Счета"]) for d in bm.debts_tbl if bm.s(d["Счета"]) not in ("Итог", "СБЕР Кредитка")]
for name in debt_names:
    seen[cp_name(name)].add("займ")
counterparties = []
cp_id = {}
for name in sorted(seen):
    kinds = seen[name]
    kind = ("person" if name in staff or "займ" in kinds else
            "buyer" if kinds == {"доход"} else "supplier" if kinds == {"расход"} else "other")
    cid = uid("counterparty", name)
    cp_id[name] = cid
    counterparties.append({
        "id": cid, "name": name, "kind": kind,
        "employeeName": name if name in staff else None,
    })

# --- Займы -----------------------------------------------------------------
loan_ops = [o for o in bm.ops if bm.s(o["Контрагент"]) != "СБЕР Кредитка"]
debts = []
debt_id = {}
for d in bm.debts_tbl:
    name = bm.s(d["Счета"])
    if name in ("Итог", "СБЕР Кредитка"):
        continue
    its = [o for o in loan_ops if bm.s(o["Контрагент"]) == name]
    start = d["Начальный остаток"] or 0
    if not its and not start:
        continue  # займа не было
    given = any("выдан" in bm.s(o["Тип операции"]).lower() for o in its)
    did = uid("debt", name)
    debt_id[name] = did
    debts.append({
        "id": did, "counterpartyId": cp_id[cp_name(name)],
        "kind": "given" if given else "received",
        "openingBalance": kop(start), "openingDate": "2025-01-01" if start else None,
    })
# Ошибочный перевод продавцу — займ выданный (решение владельца 2026-10-08).
wrong = [r for r in bm.rows if r["action"] == "долг" and "ошибоч" in r["note"].lower()]
for r in wrong:
    name = cp_name(r["who"])
    if name not in cp_id:
        cp_id[name] = uid("counterparty", name)
        counterparties.append({"id": cp_id[name], "name": name, "kind": "supplier", "employeeName": None})
    debt_id[name] = uid("debt", name)
    debts.append({"id": debt_id[name], "counterpartyId": cp_id[name], "kind": "given",
                  "openingBalance": 0, "openingDate": None})

# --- Имущество -------------------------------------------------------------
# Направления — решение владельца 2026-10-08: трактора и телеги — ключ;
# техника кормов и зерна, земля, зерносклад — растениеводство; фидлоты,
# родильная, весы, электропастух, погрузчик — животноводство; убойный цех —
# реализация; гараж, склад ГСМ, самосвал, автомобили — общехозяйственные.
ASSET_RULES = [
    ("трактор", "Т"), ("телег", "Т"),
    ("комбайн", "К"), ("пресс", "К"), ("косилк", "К"), ("грабли", "К"), ("плуг", "К"),
    ("борона", "К"), ("разбрасыват", "К"), ("опрыскиват", "К"), ("кусторез", "К"),
    ("фреза", "К"), ("зерносклад", "К"),
    ("фидлот", "С"), ("родильн", "С"), ("весы", "С"), ("электропастух", "С"),
    ("погрузчик", "С"), ("зернодробилк", "С"), ("денник", "С"), ("шкур", "С"),
    ("убойн", "Р"), ("крематор", "Р"),
    ("гараж", "О"), ("гсм", "О"), ("самосвал", "О"), ("бытовка", "О"), ("дом на колесах", "О"),
]


def asset_target(name):
    low = name.lower()
    for word, code in ASSET_RULES:
        if word in low:
            return target(code)
    raise SystemExit(f"Нет направления для имущества: {name}")


awb = openpyxl.load_workbook(ASSETS_SRC, data_only=True)
ws = awb["Расшифровка стоимости"]


def asset_table(name):
    t = ws.tables[name]
    c1, r1, c2, r2 = openpyxl.utils.range_boundaries(t.ref)
    return [[ws.cell(r, c).value for c in range(c1, c2 + 1)] for r in range(r1 + 1, r2 + 1)
            if ws.cell(r, c1 + 1).value not in (None, "")]


assets = []
asset_by_word = {}
for table_name, kind in (("Таблица1", "building"), ("Таблица2", "machinery"), ("Таблица3", "equipment")):
    for _, name, year, qty, value in asset_table(table_name):
        name = " ".join(str(name).replace("«", "").replace("»", "").split())
        if kind == "machinery" and "самосвал" in name.lower():
            kind_here = "vehicle"
        else:
            kind_here = kind
        direction, key = asset_target(name)
        aid = uid("asset", name)
        assets.append({
            "id": aid, "name": name, "kind": kind_here, "year": year, "quantity": int(qty or 1),
            "referenceValue": kop(value * 1000) if value else None,
            "valuedOn": None, "directionId": direction, "allocationKeyId": key,
            "cadastralNumber": None, "areaHa": None, "cadastralValue": None,
            "notes": "Оценка из «Имущество КФХ.xlsx»",
        })
        asset_by_word[name.lower()] = aid
# Земля: в файле «Площадь, м2» — на деле гектары (в сумме 297 га),
# «Кадастровая стоимость, тыс. руб.» — на деле рубли; «Продажная цена» —
# тысячи рублей (справочная стоимость).
for _, number, area, cadastral, price in asset_table("Таблица6"):
    direction, key = target("К")
    assets.append({
        "id": uid("asset", "land", number), "name": f"Участок {number}", "kind": "land",
        "year": None, "quantity": 1, "referenceValue": kop(price * 1000) if price else None,
        "valuedOn": None, "directionId": direction, "allocationKeyId": key,
        "cadastralNumber": str(number), "areaHa": round(float(area), 6) if area else None,
        "cadastralValue": kop(cadastral) if cadastral else None,
        "notes": "Оценка из «Имущество КФХ.xlsx»",
    })
# Добавлены владельцем 2026-10-08 (в файле их нет).
for name, value in (("Газель", 1500000), ("Нива", None)):
    aid = uid("asset", name)
    asset_by_word[name.lower()] = aid
    assets.append({
        "id": aid, "name": name, "kind": "vehicle", "year": None, "quantity": 1,
        "referenceValue": kop(value) if value else None, "valuedOn": None,
        "directionId": dir_id["О"], "allocationKeyId": None,
        "cadastralNumber": None, "areaHa": None, "cadastralValue": None,
        "notes": "Добавлено владельцем при переносе",
    })


def asset_id_like(word):
    for name, aid in asset_by_word.items():
        if word in name:
            return aid
    raise SystemExit(f"Нет имущества «{word}»")


# Привязка расходов к машине — только однозначные слова (КУН — на МТЗ-82).
ASSET_LINKS = [
    ("газел", asset_id_like("газель")),
    ("нива", asset_id_like("нива")),
    ("кун", asset_id_like("мтз-82")),
]

# --- Операции --------------------------------------------------------------
operations = []


def op(**fields):
    base = {
        "toAccountId": None, "categoryId": None, "directionId": None, "allocationKeyId": None,
        "assetId": None, "season": None, "counterpartyId": None, "employeeName": None,
        "debtId": None, "debtMove": None, "description": None, "quantity": None,
        "unit": None, "tag": None, "notes": None,
    }
    base.update(fields)
    operations.append(base)


for r in bm.rows:
    if r["action"] != "статья":
        continue
    ref = f'{r["sheet"]}!{r["row"]}'
    direction, key = (None, None)
    if r["code"] != cat_code[(r["kind"], r["new"])]:
        direction, key = target(r["code"])
    low = (r["goods"] + " " + r["note"]).lower()
    asset = next((aid for word, aid in ASSET_LINKS if word in low), None)
    who = cp_name(r["who"]) if r["who"] else None
    op(id=uid("op", ref), date=day(r["date"]), type=TYPE[r["kind"]], amount=kop(r["sum"]),
       accountId=acc_id[r["wallet"]], categoryId=cat_id[(r["kind"], r["new"])],
       directionId=direction, allocationKeyId=key, assetId=asset,
       counterpartyId=cp_id.get(who), employeeName=who if r["old"].startswith("03.") else None,
       description=text(r["goods"]), notes=text(r["note"]),
       quantity=float(r["rawqty"]) if r["kind"] == "доход" and isinstance(r["rawqty"], (int, float)) else None,
       tag=text(r["rawtag"]), sourceRef=ref)

for t in bm.table("Переводы", "Операции"):
    ref = f'Переводы!{t["_row"]}'
    op(id=uid("op", ref), date=day(t["Дата"]), type="transfer", amount=kop(t["Сумма"]),
       accountId=acc_id[bm.s(t["Счёт-отправитель"])], toAccountId=acc_id[bm.s(t["Счёт-получатель"])],
       description=text(t["Содержание операции"]), sourceRef=ref)
    assert not t["Комиссия"], "комиссия перевода — отдельный расход, в файле её нет"

MOVES = {"займ полученный": "take", "возврат полученного займа": "repay",
         "займ выданный": "lend", "возврат выданного займа": "collect"}
for o in bm.ops:
    ref = f'Долги!{o["_row"]}'
    name = bm.s(o["Контрагент"])
    wallet = acc_id[bm.s(o["Счёт зачисления/списания"])]
    if name == "СБЕР Кредитка":
        # Покупка кредиткой — перевод с кредитки (решение владельца).
        op(id=uid("op", ref), date=day(o["Дата"]), type="transfer", amount=kop(o["Сумма"]),
           accountId=acc_id["СБЕР Кредитка"], toAccountId=wallet,
           description=text(o["Комментарий"]), sourceRef=ref)
        continue
    op(id=uid("op", ref), date=day(o["Дата"]), type="debt", amount=kop(o["Сумма"]), accountId=wallet,
       debtId=debt_id[name], debtMove=MOVES[bm.s(o["Тип операции"]).lower()],
       counterpartyId=cp_id[cp_name(name)], description=text(o["Комментарий"]), sourceRef=ref)

for r in wrong:
    ref = f'{r["sheet"]}!{r["row"]}'
    name = cp_name(r["who"])
    op(id=uid("op", ref), date=day(r["date"]), type="debt", amount=kop(r["sum"]),
       accountId=acc_id[r["wallet"]], debtId=debt_id[name], debtMove="lend",
       counterpartyId=cp_id[name], description=text(r["goods"]), notes=text(r["note"]), sourceRef=ref)

# Строки займов в «Доход»/«Расход» — дубли реестра «Долги» (сверено в 8.0).
skipped = [f'{r["sheet"]}!{r["row"]}' for r in bm.rows
           if r["action"] == "долг" and "ошибоч" not in r["note"].lower()]

# --- Выплаты ЗП → расход (решения владельца 2026-10-08) --------------------
payroll = {
    "methodAccounts": {"cash": acc_id["Касса нал"], "card": acc_id["Карта ГКФХ"], "transfer": acc_id["СБЕР"],
                       # Выплата из подотчёта (Нал Ольга) — импорт «Финансов 25».
                       "accountable": acc_id["Нал Ольга"]},
    "defaultAccountId": acc_id["Касса нал"],
    "positionCategories": {
        "механизатор": cat_id[("расход", "Механизаторы")],
        "мезанизатор": cat_id[("расход", "Механизаторы")],  # так в табеле
        "глава хозяйства": cat_id[("расход", "Глава хозяйства")],
        "главай хозяйства": cat_id[("расход", "Глава хозяйства")],  # так в табеле
        "маркетолог": cat_id[("расход", "Продавец, сбытовик")],
        "рабочий": cat_id[("расход", "Рабочие фермы")],
        # Должности работников 2025 года (по статьям их ЗП в «Финансах 25»).
        "главный бухгалтер": cat_id[("расход", "Главный бухгалтер")],
        "управляющий": cat_id[("расход", "Управляющий")],
        "продавец": cat_id[("расход", "Продавец, сбытовик")],
        "мясник": cat_id[("расход", "Мясник")],
    },
    "defaultCategoryId": cat_id[("расход", "Рабочие фермы")],
}

# --- «Табель 25» (решения владельца 2026-10-08) -----------------------------
# Работники, которых нет в табеле программы (по «Фамилия И.»), добавляются
# уволенными: приём — первый отработанный день 2025–2026, увольнение —
# следующий за последним; ставка базы и поля — тариф месяца последнего дня
# (нет — ближайшего); должность — по статье их ЗП в «Финансах 25».
# Дни 2025 года: 1 и 0,5 — как есть, больше 1 — один день, остальные доли —
# 0,5; место — база (в Excel его нет). Начислено за месяц — «Сумма» Excel:
# разницу с расчётом программы (дни сверх, доли, надбавки, премии, вычеты,
# другой тариф) импортёр записывает корректировкой месяца. Выплаты — из
# расходов ЗП «Финансов 25» (ниже), «Выдал» табеля — только для сравнения.
MONTHS = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "июн": 6,
          "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12}
TABEL_YEAR = 2025


def number(value):
    if isinstance(value, str):  # «1 ,5» в августе
        try:
            value = float(value.replace(" ", "").replace(",", "."))
        except ValueError:
            return 0
    return value if isinstance(value, (int, float)) else 0


def column(head, title):
    return next(i for i, v in enumerate(head) if isinstance(v, str) and v.strip().startswith(title))


people = {}
months_data = []  # (имя, год, месяц, {день: значение}, строка итогов)
tabel = openpyxl.load_workbook(TABEL_SRC, data_only=True, read_only=True)
for ws in tabel.worksheets:
    parts = ws.title.split()
    if len(parts) != 2 or parts[0] not in MONTHS:
        continue  # шаблон, итоги года, пустые листы
    year, month = 2000 + int(parts[1]), MONTHS[parts[0]]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[4]
    day_cols = [i for i, v in enumerate(head) if isinstance(v, int)]
    cols = {k: column(head, t) for k, t in (
        ("days", "Итого отработано"), ("tariff", "Тариф"), ("bonus", "Надбавки"),
        ("premium", "Премия"), ("sum", "Сумма"), ("paid", "Выдал"))}
    for r in rows[5:]:
        name = text(r[1]) if len(r) > 1 else None
        if not name or name == "Итого" or not isinstance(r[0], int):
            continue
        p = people.setdefault(name, {"first": None, "last": None, "rate": {}})
        days = {}
        for i in day_cols:
            v = number(r[i])
            if v > 0:
                d = datetime.date(year, month, head[i])
                days[d] = v
                p["first"] = p["first"] or d
                p["last"] = d
        if r[cols["tariff"]]:
            p["rate"][(year, month)] = r[cols["tariff"]]
        months_data.append((name, year, month, days, {k: number(r[i]) for k, i in cols.items()}))


def short(name):
    """«Иванов И. П.» → («иванов», «и») — как matchEmployee в импортёре."""
    w = name.replace(".", " ").split()
    return w[0].lower().replace("ё", "е"), (w[1][0].lower() if len(w) > 1 else None)


# Выплаты ЗП «Финансов 25» → выплаты табеля: расход из них программа строит
# сама (кошелёк — по способу, статья — по должности), поэтому из операций
# они убираются — иначе расход задвоится.
METHOD = {"Касса нал": "cash", "Карта ГКФХ": "card", "СБЕР": "transfer"}
acc_by_id = {a["id"]: a for a in accounts}
accountable = [a for a in accounts if a["kind"] == "accountable"]
assert len(accountable) == 1, "способ «подотчёт» — ровно один кошелёк"
cat_name = {c["id"]: c["name"] for c in categories}
payments = []
position_of = {}
for o in operations:
    if not o["employeeName"]:
        continue
    acc = acc_by_id[o["accountId"]]
    method = "accountable" if acc["kind"] == "accountable" else METHOD[acc["name"]]
    payments.append({
        "id": uid("payment", o["sourceRef"]), "employeeName": o["employeeName"],
        "date": o["date"], "amount": o["amount"], "method": method,
        "notes": "; ".join(t for t in (o["description"], o["notes"]) if t) or None,
        "sourceRef": o["sourceRef"],
    })
    position_of.setdefault(short(o["employeeName"]), set()).add(cat_name[o["categoryId"]])
operations = [o for o in operations if not o["employeeName"]]
# Должность → статья ЗП (у каждого работника статья одна).
CAT_POSITION = {"Механизаторы": "Механизатор", "Рабочие фермы": "Рабочий",
                "Глава хозяйства": "Глава хозяйства", "Главный бухгалтер": "Главный бухгалтер",
                "Управляющий": "Управляющий", "Продавец, сбытовик": "Продавец", "Мясник": "Мясник"}
assert all(len(c) == 1 for c in position_of.values()), position_of

employees = []
for name, p in people.items():
    if not p["first"]:
        continue
    last = p["last"]
    cats = position_of.get(short(name))
    employees.append({
        "id": uid("employee", name), "name": name,
        "position": CAT_POSITION[next(iter(cats))] if cats else "Рабочий",
        "hireDate": day(p["first"]), "dismissalDate": day(last + datetime.timedelta(days=1)),
        # В месяце без тарифа (Лопатин А., январь — разовая сумма) — ближайший.
        "rate": p["rate"].get((last.year, last.month))
        or min(p["rate"].items(), key=lambda kv: abs((kv[0][0] - last.year) * 12 + kv[0][1] - last.month))[1],
    })

timesheet = []
accruals = []
for name, year, month, days, totals in months_data:
    if year != TABEL_YEAR:
        continue
    for d, v in sorted(days.items()):
        timesheet.append({
            "id": uid("timesheet", name, day(d)), "employeeName": name, "date": day(d),
            "days": 1 if v >= 1 else 0.5, "excel": v,
        })
    if days or totals["sum"] or totals["paid"]:
        accruals.append({
            "id": uid("adjustment", name, f"{year}-{month:02}"), "employeeName": name,
            "year": year, "month": month, "excelDays": totals["days"], "tariff": totals["tariff"],
            "bonus": totals["bonus"], "premium": totals["premium"],
            "accrued": totals["sum"], "paid": totals["paid"],
        })

tabel_keys = {short(e["name"]) for e in employees}
finance_people = {x["employeeName"] for k in ("accounts", "counterparties", "payments")
                  for x in globals()[k] if x.get("employeeName")}
missing = sorted(n for n in finance_people if short(n) not in tabel_keys)
assert not missing, f"работники «Финансов 25» не найдены в «Табеле 25»: {missing}"

# --- Контрольные итоги: из Excel, независимо от операций --------------------
control = collections.defaultdict(int)


def add(key, value):
    control[key] += value


for r in bm.rows:
    if r["action"] == "статья" or "ошибоч" in r["note"].lower():
        add(f'{r["sheet"]}|{r["date"]:%Y-%m}', kop(r["sum"]))
        add(f'{r["sheet"]}|всего', kop(r["sum"]))
for t in bm.table("Переводы", "Операции"):
    add("Переводы|всего", kop(t["Сумма"]))
for o in bm.ops:
    add("Долги|всего", kop(o["Сумма"]))

out = {
    "source": "Финансы 25.xlsm, Имущество КФХ.xlsx, Табель 25.xlsx",
    "generated": datetime.datetime.now().isoformat(timespec="seconds"),
    "directions": directions, "allocationKeys": keys, "categories": categories,
    "accounts": accounts, "counterparties": counterparties, "debts": debts, "assets": assets,
    "operations": operations, "payrollMapping": payroll, "employees": employees, "timesheet": timesheet,
    "accruals": accruals, "payments": payments, "skipped": skipped,
    "control": dict(sorted(control.items())),
}
assert len({o["id"] for o in operations}) == len(operations), "повтор uuid операции"
path = os.path.join(bm.PRIVATE, "finance25_import.json")
with open(path, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1)
kinds = collections.Counter(o["type"] for o in operations)
print("выгрузка:", path)
print("операций:", dict(kinds), "всего", len(operations), "| пропущено дублей:", len(skipped))
print("статей", len(categories), "кошельков", len(accounts), "контрагентов", len(counterparties),
      "займов", len(debts), "имущества", len(assets))
print("с машиной:", sum(1 for o in operations if o["assetId"]),
      "| своё направление:", sum(1 for o in operations if o["directionId"] or o["allocationKeyId"]))
print("«Табель 25»: работников", len(employees), "| дней", len(timesheet),
      "| месяцев начислений", len(accruals), "| выплат из «Финансов 25»", len(payments))
